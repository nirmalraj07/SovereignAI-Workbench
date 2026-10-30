"""
Text extraction for ingest — turns uploaded PDF/DOCX/CSV/XLSX/TXT bytes into
plain text (with a page map for PDFs) before chunking/embedding in
rag/pipeline.py. Runs fully offline (pdfplumber, python-docx, openpyxl —
no external service calls).
"""
from __future__ import annotations
import csv
import io


def extract_text(raw: bytes, filename: str) -> tuple[str, list[tuple[str, int]] | None]:
    """Returns (full_text, page_map). page_map is [(text, page_number)] for
    paginated formats (PDF), or None when the format has no page concept."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "pdf":
        return _extract_pdf(raw)
    if ext == "docx":
        return _extract_docx(raw), None
    if ext in ("csv",):
        return _extract_csv(raw), None
    if ext in ("xlsx", "xls"):
        return _extract_xlsx(raw), None
    # txt and anything unrecognized: treat as plain text
    return raw.decode(errors="ignore"), None


def _extract_pdf(raw: bytes) -> tuple[str, list[tuple[str, int]]]:
    import pdfplumber

    page_map: list[tuple[str, int]] = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            page_map.append((text, i))
            # Table extraction: flatten detected tables into text too, so
            # tabular SOP/inspection data is searchable, not just prose.
            for table in page.extract_tables():
                rows_text = "\n".join(" | ".join(cell or "" for cell in row) for row in table)
                page_map.append((rows_text, i))
    full_text = " ".join(t for t, _ in page_map)
    return full_text, page_map


def _extract_docx(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_csv(raw: bytes) -> str:
    text = raw.decode(errors="ignore")
    reader = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(row) for row in reader)


def _extract_xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            parts.append(" | ".join(str(c) if c is not None else "" for c in row))
    return "\n".join(parts)
